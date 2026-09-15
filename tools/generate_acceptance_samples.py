from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUTPUT = Path(__file__).resolve().parents[1] / "acceptance" / "samples"

SAMPLES = (
    (
        "01-人事管理制度.docx",
        "员工培训管理制度",
        "人事制度",
        (
            ("适用范围", "本制度适用于单位内部员工的岗前培训和年度继续教育。"),
            ("培训安排", "综合部门每年制定培训计划，记录培训主题、参加人员和完成情况。"),
            ("资料管理", "培训资料由经办人员整理后归档，供内部查询和引用。"),
        ),
    ),
    (
        "02-财务报销办法.docx",
        "差旅费用报销办法",
        "财务制度",
        (
            ("报销范围", "差旅费包括交通费、住宿费和按规定发放的出差补助。"),
            ("提交要求", "报销人应提交审批记录和有效票据，并注明出差事由。"),
            ("审核流程", "部门负责人确认后交财务复核，资料完整后办理报销。"),
        ),
    ),
    (
        "03-项目会议纪要.docx",
        "资料管理项目会议纪要",
        "项目资料",
        (
            ("会议议题", "讨论离线文档导入、正文搜索、分类管理和诊断日志方案。"),
            ("会议决定", "文件导入资料库时采用复制方式，一个文档只属于一个分类。"),
            ("后续事项", "使用虚构材料开展验收，不在诊断包中记录文件名和正文。"),
        ),
    ),
    (
        "04-普通工作通知.docx",
        "办公区域调整通知",
        "综合通知",
        (
            ("调整时间", "办公区域调整工作定于本月最后一个工作日进行。"),
            ("注意事项", "请提前整理个人物品，公共资料统一装箱并做好编号。"),
            ("联系事项", "如有设备搬运需求，请与综合部门联系。"),
        ),
    ),
)


def set_run_font(run, size: int, *, bold: bool = False) -> None:
    run.font.name = "Noto Sans CJK SC"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Noto Sans CJK SC")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = None


def create_sample(filename: str, title: str, category: str, sections) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)

    title_paragraph = document.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title_paragraph.add_run(title), 20, bold=True)

    intro = document.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(intro.add_run(f"模拟验收资料  建议分类 {category}"), 10)

    for heading, body in sections:
        heading_paragraph = document.add_paragraph(style="Heading 1")
        set_run_font(heading_paragraph.add_run(heading), 14, bold=True)
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
        set_run_font(paragraph.add_run(body), 11)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("仅用于文澜资料库离线验收"), 9)

    OUTPUT.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT / filename)


def main() -> None:
    for sample in SAMPLES:
        create_sample(*sample)
    for index in range(5, 65):
        create_stress_sample(index)


def create_stress_sample(index: int) -> None:
    topics = ("人事培训", "财务报销", "项目管理", "档案保管", "安全检查", "综合通知")
    topic = topics[index % len(topics)]
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    if index % 10 == 0:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run(f"{topic}模拟文档 {index:02d}"), 20, bold=True)
    intro = document.add_paragraph()
    set_run_font(intro.add_run("本文档仅使用虚构内容，用于验证离线导入、预览和搜索。"), 11)

    paragraph_count = 2 + index % 9
    for chapter in range(1, paragraph_count + 1):
        heading = document.add_paragraph(style="Heading 1" if chapter < 4 else "Heading 2")
        set_run_font(
            heading.add_run(f"第{chapter}部分  测试事项"),
            14 if chapter < 4 else 12,
            bold=True,
        )
        body = document.add_paragraph()
        body.paragraph_format.line_spacing = 1.5
        body.paragraph_format.first_line_indent = Cm(0.74)
        marker = f"SIM-{index:02d}-{chapter:02d}"
        text = f"{topic}的模拟记录编号为 {marker}。所有单位、人员、金额和日期均为虚构数据。"
        if index % 5 == 0:
            text += " 包含中英文 English、数字 2026.09 与符号（甲）。"
        set_run_font(body.add_run(text), 11)

    if index % 3 == 0:
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells, ("序号", "事项", "日期", "状态"), strict=True):
            set_run_font(cell.paragraphs[0].add_run(value), 10, bold=True)
        for row_index in range(1, 3 + index % 8):
            cells = table.add_row().cells
            values = (str(row_index), f"模拟任务 {row_index}", f"2026-09-{row_index:02d}", "已完成")
            for cell, value in zip(cells, values, strict=True):
                set_run_font(cell.paragraphs[0].add_run(value), 10)

    if index % 4 == 0:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(header.add_run("模拟资料页眉"), 9)
    if index % 7 == 0:
        document.add_page_break()
        heading = document.add_paragraph(style="Heading 1")
        set_run_font(heading.add_run("附录"), 14, bold=True)
        set_run_font(document.add_paragraph().add_run("用于验证分页和长文档预览。"), 11)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(footer.add_run("仅用于离线功能验收"), 9)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT / f"{index:02d}-{topic}模拟文档.docx")


if __name__ == "__main__":
    main()
